import json
import os
import sys
from docx import Document

src, out_dir = sys.argv[1], sys.argv[2]
os.makedirs(out_dir, exist_ok=True)
doc = Document(src)
items = []
for p in doc.paragraphs:
    text = p.text.strip()
    if text:
        items.append({"type": "paragraph", "style": p.style.name if p.style else "", "text": text})
for ti, table in enumerate(doc.tables, 1):
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    items.append({"type": "table", "index": ti, "rows": rows})
with open(os.path.join(out_dir, "content.json"), "w", encoding="utf-8") as f:
    json.dump(items, f, ensure_ascii=False, indent=2)
with open(os.path.join(out_dir, "content.txt"), "w", encoding="utf-8") as f:
    for item in items:
        if item["type"] == "paragraph":
            f.write(f'[{item["style"]}] {item["text"]}\n')
        else:
            f.write(f'\n[TABLE {item["index"]}]\n')
            for row in item["rows"]:
                f.write(" | ".join(row) + "\n")
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        part = rel.target_part
        name = os.path.basename(part.partname)
        with open(os.path.join(out_dir, name), "wb") as f:
            f.write(part.blob)
print(json.dumps({"paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "out": out_dir}, ensure_ascii=False))
